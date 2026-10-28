from dataclasses import replace
from datetime import date
from openpyxl import load_workbook
import pytest
from trazabilidad.template import create_template
from trazabilidad.config import Config
from trazabilidad.excel_reader import read_plan


def basic(path, own=0, third=0):
    create_template(path,example=True)
    book = load_workbook(path)
    for name, count in [('Plantaciones Propias',own), ('Plantaciones de Terceros',third)]:
        ws=book[name]
        for ix in range(6,6+count):
            ws[f'C{ix}'] = 'Empresa Ñandú Ficticia'
            ws[f'D{ix}'] = f'Plantación TEST {ix}'
            ws[f'E{ix}'] = f'00{ix}'
            ws[f'F{ix}'] = 'NO'
            ws[f'G{ix}'] = 7
            ws[f'H{ix}'] = 3.4869034155
            ws[f'I{ix}'] = -73.5440480795
            ws[f'N{ix}'] = 2019
            ws[f'O{ix}'] = 15.578
            ws[f'Q{ix}'] = 'NO'
    book.save(path)


def test_large_workbook(tmp_path):
    p=tmp_path/'Archivo rural en español.xlsx'
    basic(p,73,275)
    plan=read_plan(p,Config())
    assert not plan.blockers,plan.blockers
    assert len(plan.own) == 73
    assert len(plan.third) == 275
    assert plan.third[-1].excel_row == 280
    assert plan.own[0].values['fruto_tm'] == '15.58'
    assert plan.own[0].values['latitud'] == '3.486903'
    assert plan.own[0].values['longitud'] == '-73.544048'


def test_id_duplicates_and_empty_vs_zero(tmp_path):
    p=tmp_path/'dups.xlsx';basic(p,2,0)
    b=load_workbook(p); w=b['Plantaciones Propias']
    w['E7']='006';w['O6']=0;w['O7']=None;w['Q6']='SI'
    b.save(p)
    plan=read_plan(p,Config())
    assert not plan.blockers
    assert plan.own[0].values['fruto_tm']=='0.00'
    assert plan.own[1].values['fruto_tm'] is None
    assert any('repetido' in issue.message for issue in plan.issues)
    assert any('Adjuntar' in issue.message for issue in plan.issues)


def test_multiple_year_and_override(tmp_path):
    p=tmp_path/'years.xlsx'; basic(p,1,0)
    b=load_workbook(p);b['Plantaciones Propias']['N6']='2018;2023;2023';b.save(p)
    assert read_plan(p,Config()).blockers
    accepted=read_plan(p,replace(Config(),year_override_own=2019))
    assert not accepted.blockers
    assert accepted.own[0].values['anio_certificacion']=='2019'
    assert accepted.decisions


def test_zero_year_requires_decision(tmp_path):
    p=tmp_path/'zero.xlsx'; basic(p,0,1)
    b=load_workbook(p);b['Plantaciones de Terceros']['N6']=0;b.save(p)
    assert read_plan(p,Config()).blockers
    plan=read_plan(p,replace(Config(),third_zero_year='empty'))
    assert plan.third[0].values['anio_certificacion'] is None
    assert not plan.blockers


def test_formula_without_cache(tmp_path):
    p=tmp_path/'formula.xlsx';basic(p,1,0)
    b=load_workbook(p);b['Plantaciones Propias']['O6']='=A1+1';b.save(p)
    plan=read_plan(p,Config())
    assert any('cacheado' in i.message for i in plan.blockers)


def test_bad_heading_blocks(tmp_path):
    p=tmp_path/'heading.xlsx';basic(p)
    b=load_workbook(p);b['Plantaciones Propias']['O5']='Columna inesperada';b.save(p)
    assert any(i.field=='O' for i in read_plan(p,Config()).blockers)


def test_original_preserved(tmp_path):
    p=tmp_path/'original.xlsx';basic(p)
    original=p.read_bytes()
    read_plan(p,Config())
    assert original==p.read_bytes()
