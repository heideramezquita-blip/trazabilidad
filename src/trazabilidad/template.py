"""Genera plantilla editable compatible con v1; nunca modifica Excel original."""
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from .models import SHEETS, HEADERS_MILL, HEADERS_PLANT

BLUE = '183153'
BLUE_LIGHT = 'DDEAF6'
EDIT = 'FFFFFF'
AUTO = 'EBEEF1'
AUTOMATIC_MILL = {'A', 'K', 'L', 'M'}
AUTOMATIC_PLANT = {'A', 'B', 'J', 'K', 'L', 'P'}


def create_template(path, example=False):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        raise FileExistsError(f'No se sobrescribe el archivo: {target}')
    wb = Workbook()
    main = wb.active
    main.title = SHEETS[0]
    own = wb.create_sheet(SHEETS[1]); third = wb.create_sheet(SHEETS[2])
    main['B2'] = 'REPORTE DE TRAZABILIDAD · PLANTILLA LOCAL V2'
    main['B2'].font = Font(size=17, bold=True, color=BLUE)
    for pos, label in [('B5', 'Periodo (C5)'), ('B6', 'Fecha de referencia Excel (C6)'), ('B7', 'Diligenciador (C7)')]:
        main[pos] = label
        main[pos].font = Font(bold=True, color=BLUE)
    main['C6'].number_format = 'dd/mm/yyyy'
    main['B9'] = 'Los campos grises son automáticos del sitio y no se transfieren desde Excel.'
    main['B9'].font = Font(italic=True, color='52657C')
    main['B14'] = 'INSTRUCCIONES: solo la fila 11 contiene el registro del molino.'
    main['B15'] = 'Fecha C6: referencia del archivo; la fecha efectiva de envío es la del sitio.'
    main['B16'] = 'Formato visible ≠ redondeo real: el cargador normaliza TM, área y coordenadas.'
    main['B17'] = 'Los datos no se envían automáticamente; revisa la página antes de enviar.'
    if example:
        main['C5'] = 'Q3 - 2026'
        main['C6'] = date(2026, 9, 30)
        main['C6'].number_format = 'dd/mm/yyyy'
        main['C7'] = 'Persona Ficticia'
    for ws, row, columns, start, auto in [(main, 10, HEADERS_MILL, 11, AUTOMATIC_MILL),
                                           (own, 5, HEADERS_PLANT, 6, AUTOMATIC_PLANT),
                                           (third, 5, HEADERS_PLANT, 6, AUTOMATIC_PLANT)]:
        ws.freeze_panes = f'C{start}'
        ws.sheet_view.showGridLines = False
        for col, header in columns.items():
            head = ws[f'{col}{row}']; head.value = header
            head.fill = PatternFill('solid', fgColor=BLUE)
            head.font = Font(color='FFFFFF', bold=True, size=10)
            head.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.column_dimensions[col].width = 19 if col not in ('C','D','E','G','M') else 25
        ws.row_dimensions[row].height = 38
        for i in range(start, start+350):
            for col in columns:
                cell = ws[f'{col}{i}']
                cell.fill = PatternFill('solid', fgColor=AUTO if col in auto else EDIT)
                cell.alignment = Alignment(vertical='center')
                if col in ('H', 'I'):
                    cell.number_format = '0.000000'
                if col in (('P', 'Q') if ws == main else ('G', 'O')):
                    cell.number_format = '0.00'
                if col == ('B' if ws == main else 'E'):
                    cell.number_format = '@'
        ws.auto_filter.ref = f'A{row}:Q{start+349}'
        rspo = DataValidation(type='list', formula1='"NO,MB,SG,IP"', allow_blank=True)
        rspo.error = 'Solo NO, MB, SG o IP'; rspo.errorTitle = 'Certificación RSPO'; rspo.showErrorMessage = True
        ws.add_data_validation(rspo)
        rspo.add(f'F{start}:F{start+349}')
        if ws != main:
            poly = DataValidation(type='list', formula1='"SI,NO"', allow_blank=True)
            poly.error = 'Selecciona SI o NO'; poly.showErrorMessage = True
            ws.add_data_validation(poly); poly.add(f'Q{start}:Q{start+349}')
            year = DataValidation(type='whole', operator='between', formula1='1000', formula2='9999', allow_blank=True)
            year.error = 'Año único de cuatro cifras'; year.showErrorMessage = True
            ws.add_data_validation(year); year.add(f'N{start}:N{start+349}')
    if example:
        values = {'B11':'UM-FICT-01', 'C11':'Organización Ficticia', 'D11':'Molino Ejemplo S.A.S.',
                  'E11':'Molino Demostración', 'F11':'MB', 'G11':'Ejemplo',
                  'H11':3.486903, 'I11':-73.544048, 'J11':'Vía Ficticia Km 1',
                  'P11':102.83419, 'Q11':0}
        for key, value in values.items(): main[key] = value
        for ws, count, prefix in [(own, 3, 'PROPIA'), (third, 4, 'TERCERA')]:
            for n in range(count):
                r = n + 6
                for col, value in {'C': 'Grupo Ejemplo', 'D': f'Plantación {prefix} {n+1}',
                                   'E': f'00{n+1}', 'F':'NO', 'G': 7.56, 'H':3.486903,
                                   'I': -73.544048, 'M': 'Ficticia', 'N': 2019,
                                   'O': 15.578, 'Q':'NO'}.items():
                    ws[f'{col}{r}'] = value
                ws[f'E{r}'].number_format = '@'
    wb.save(target)
    return target
