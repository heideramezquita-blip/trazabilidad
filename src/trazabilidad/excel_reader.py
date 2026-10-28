"""Lee la plantilla v1/v2 sin escribir sobre el libro original."""
from __future__ import annotations
from datetime import datetime, date
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import re
import unicodedata
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from .models import (
    Plan, Record, Issue, SHEETS, MILL_FIELDS, PLANT_FIELDS,
    HEADERS_MILL, HEADERS_PLANT, AUTO_FIELDS, TEXT_FIELDS, DECIMAL_FIELDS,
    NUMERIC_OTHER, SELECT_FIELDS,
)
from .normalize import normalize_decimal, ValueErrorLocal


def normalize_heading(value):
    s = unicodedata.normalize('NFKD', str(value or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c)).casefold()
    return re.sub(r'[^a-z0-9]+', '', s)


def as_text(value, cell=None):
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        return s or None
    if isinstance(value, int) and cell is not None and re.fullmatch(r'0{2,}', cell.number_format or ''):
        return str(value).zfill(len(cell.number_format))
    return str(value)


def _cell_value(ws, cached, pos, issues, sheet, row, field):
    c = ws[pos]
    value = cached[pos].value
    if c.data_type == 'f' and value is None:
        issues.append(Issue('ERROR', sheet, row, field,
                            f'Fórmula {pos} sin resultado cacheado: recalcular y guardar en Excel.'))
    return value


def parse_record(ws, cached, rownum, mapping, plan, config):
    values = {}
    for col, field in mapping.items():
        value = _cell_value(ws, cached, f'{col}{rownum}', plan.issues, ws.title, rownum, field)
        if field in AUTO_FIELDS:
            continue
        try:
            if field in TEXT_FIELDS:
                values[field] = as_text(value, ws[f'{col}{rownum}'])
            elif field in DECIMAL_FIELDS:
                values[field] = normalize_decimal(value, DECIMAL_FIELDS[field], config.decimal_separator)
            elif field in NUMERIC_OTHER:
                values[field] = normalize_decimal(value, None, config.decimal_separator)
            elif field in SELECT_FIELDS:
                values[field] = as_text(value).upper() if as_text(value) else None
                if values[field] and values[field] not in ({'NO', 'MB', 'SG', 'IP'} if field == 'rspo' else {'SI', 'NO'}):
                    raise ValueErrorLocal(f'Opción no admitida: {values[field]}.')
            elif field == 'anio_certificacion':
                values[field] = year_value(value, ws.title, plan, config, rownum)
        except (ValueErrorLocal, ValueError) as exc:
            values[field] = None
            plan.issues.append(Issue('ERROR', ws.title, rownum, field, str(exc)))

    for field, lo, hi in [('latitud', -90, 90), ('longitud', -180, 180),
                          ('pct_propias', 0, 100), ('pct_terceros', 0, 100)]:
        if values.get(field) is not None and not Decimal(lo) <= Decimal(values[field]) <= Decimal(hi):
            plan.issues.append(Issue('ERROR', ws.title, rownum, field, f'Fuera del rango {lo} a {hi}.'))
    for field in ('fruto_tm', 'cpo_tm', 'pko_tm', 'area_ha'):
        if values.get(field) is not None and Decimal(values[field]) < 0:
            plan.issues.append(Issue('ERROR', ws.title, rownum, field, 'No puede ser negativo.'))
    if mapping == MILL_FIELDS:
        for field in ('uml_id', 'molino'):
            if not values.get(field):
                plan.issues.append(Issue('ERROR', ws.title, rownum, field, 'Identificación local obligatoria.'))
        if values.get('pct_propias') and values.get('pct_terceros'):
            total = Decimal(values['pct_propias']) + Decimal(values['pct_terceros'])
            if total != 100:
                plan.issues.append(Issue('AVISO', ws.title, rownum, 'porcentajes', f'Suman {total}, no 100.'))
    else:
        for field in ('plantacion', 'id_plantacion'):
            if not values.get(field):
                plan.issues.append(Issue('ERROR', ws.title, rownum, field, 'Identificación local obligatoria.'))
        if values.get('poligonos') == 'SI':
            plan.issues.append(Issue('AVISO', ws.title, rownum, 'poligonos', 'Adjuntar polígonos manualmente.'))
    return Record(ws.title, rownum, values, original_tm=_original_tm(ws, cached, rownum, config, plan))


def _original_tm(ws, cached, row, config, plan):
    col = 'P' if ws.title == 'Panel Principal' else 'O'
    # El molino tiene dos volúmenes; totalizar por separado en el plan.
    if ws.title == 'Panel Principal':
        return None
    value = cached[f'{col}{row}'].value
    if value is None or str(value).strip() == '':
        return None
    try:
        from .normalize import decimal_from_excel
        return str(decimal_from_excel(value, config.decimal_separator))
    except ValueErrorLocal:
        return None


def year_value(value, sheet, plan, config, row):
    text = as_text(value)
    if text is None:
        return None
    if ';' in text or ',' in text:
        if sheet == 'Plantaciones Propias' and config.year_override_own is not None:
            plan.decisions.append(f'{sheet}, fila {row}: año múltiple {text!r} reemplazado por {config.year_override_own} (decisión configurada).')
            return str(config.year_override_own)
        raise ValueErrorLocal(f'Año múltiple {text!r}; decidir un único año de cuatro cifras.')
    if text == '0':
        if sheet != 'Plantaciones de Terceros':
            raise ValueErrorLocal('Cero no es año calendario válido.')
        action = config.third_zero_year
        if action == 'empty':
            plan.decisions.append(f'{sheet}, fila {row}: año 0 reemplazado por vacío por decisión explícita.')
            return None
        if action == 'preserve':
            plan.decisions.append(f'{sheet}, fila {row}: año 0 conservado como marcador, no año calendario.')
            return '0'
        raise ValueErrorLocal('Año 0: configurar third_zero_year como empty o preserve bajo decisión humana.')
    if not re.fullmatch(r'[0-9]{4}', text):
        raise ValueErrorLocal(f'Año inválido {text!r}; se requieren cuatro cifras.')
    return text


def _validate_headers(ws, actual_row, expected, plan):
    for col, label in expected.items():
        if col == 'A':
            continue
        v = ws[f'{col}{actual_row}'].value
        if normalize_heading(v) != normalize_heading(label):
            plan.issues.append(Issue('ERROR', ws.title, actual_row, col,
                                     f'Encabezado esperado {label!r}; recibido {str(v)!r}. No asumir columnas.'))


def _is_data_row(ws, row):
    # A: numeración automática; B y otros campos auto se excluyen para detectar filas reales.
    excluded = {'A', 'B', 'J', 'K', 'L', 'P'}
    return any(ws[f'{col}{row}'].value is not None and str(ws[f'{col}{row}'].value).strip()
               for col in 'CDEFGHIMNOQ' if col not in excluded)


def read_plan(path, config):
    path = Path(path).expanduser().resolve()
    if not path.exists() or path.suffix.lower() != '.xlsx':
        raise ValueError('Selecciona un .xlsx existente (no .xlsm).')
    content_hash = sha256(path.read_bytes()).hexdigest()
    plan = Plan(str(path), content_hash)
    wb = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, read_only=False, keep_links=False)
    for sheet in SHEETS:
        if sheet not in wb.sheetnames:
            plan.issues.append(Issue('ERROR', sheet, 0, '', 'Hoja obligatoria ausente.'))
    if plan.blockers:
        return plan
    main = wb['Panel Principal']
    cache_main = cached['Panel Principal']
    _validate_headers(main, 10, HEADERS_MILL, plan)
    for sheet in SHEETS[1:]:
        _validate_headers(wb[sheet], 5, HEADERS_PLANT, plan)
    for field, pos in [('periodo', 'C5'), ('fecha_archivo', 'C6'), ('diligenciador', 'C7')]:
        val = _cell_value(main, cache_main, pos, plan.issues, main.title, int(pos[1:]), field)
        if field == 'periodo':
            plan.period = as_text(val)
        if field == 'diligenciador':
            plan.operator = as_text(val)
        if field == 'fecha_archivo':
            if isinstance(val, (date, datetime)):
                plan.file_date = val.date().isoformat() if isinstance(val, datetime) else val.isoformat()
            elif isinstance(val, (int, float)) and main[pos].is_date:
                plan.file_date = from_excel(val, wb.epoch).date().isoformat()
            else:
                plan.file_date = as_text(val)
    if not plan.operator:
        plan.issues.append(Issue('ERROR', main.title, 7, 'diligenciador', 'Nombre obligatorio.'))
    if not plan.period:
        plan.issues.append(Issue('ERROR', main.title, 5, 'periodo', 'Periodo sin diligenciar.'))
    plan.mill = parse_record(main, cache_main, 11, MILL_FIELDS, plan, config)
    for extra in (12, 13):
        if any(main[f'{col}{extra}'].value is not None and str(main[f'{col}{extra}'].value).strip()
               for col in MILL_FIELDS):
            plan.issues.append(Issue('ERROR', main.title, extra, 'molino',
                                     'Se esperaba un solo registro de molino en la fila 11.'))
    for sheet, target in [('Plantaciones Propias', plan.own), ('Plantaciones de Terceros', plan.third)]:
        ws, ws_cache = wb[sheet], cached[sheet]
        ids, originals = {}, {}
        for row in range(6, ws.max_row + 1):
            if not _is_data_row(ws, row):
                continue
            record = parse_record(ws, ws_cache, row, PLANT_FIELDS, plan, config)
            target.append(record)
            ident = record.values.get('id_plantacion')
            if ident:
                if ident in ids:
                    plan.issues.append(Issue('AVISO', sheet, row, 'id_plantacion',
                                             f'ID repetido con fila {ids[ident]}; ambos se mantienen.'))
                else:
                    ids[ident] = row
            fingerprint = tuple(sorted(record.values.items()))
            if fingerprint in originals:
                plan.issues.append(Issue('AVISO', sheet, row, 'fila',
                                         f'Fila exactamente duplicada de {originals[fingerprint]}; no eliminada.'))
            else:
                originals[fingerprint] = row
    for key, records in [('propias', plan.own), ('terceros', plan.third)]:
        orig = sum((Decimal(r.original_tm) for r in records if r.original_tm is not None), Decimal('0'))
        normalized = sum((Decimal(r.values['fruto_tm']) for r in records if r.values.get('fruto_tm') is not None), Decimal('0'))
        plan.totals[key] = {
            'registros': len(records), 'tm_original': format(orig, 'f'),
            'tm_normalizada': format(normalized, '.2f'),
            'diferencia': format(normalized - orig, 'f'),
            'tm_original_redondeada_al_final': format(orig.quantize(Decimal('.01')), '.2f'),
        }
    return plan
